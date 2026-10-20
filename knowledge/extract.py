"""
knowledge/extract.py
Text extraction and cleaning for PDF, DOCX, TXT and MD files.

Usage (PowerShell):
    python knowledge/extract.py --source data/raw --out data/processed

What it does:
  - Extracts raw text from each supported file
  - Normalises whitespace and removes non-printable characters
  - Writes one .txt per source file to the output folder
  - Skips files that have already been extracted (use --force to re-run)
"""
import argparse
import re
import sys
import unicodedata
from pathlib import Path
from typing import Optional, Tuple

from config import DATA_RAW, DATA_PROCESSED

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".rtf"}


# ── Extraction ───────────────────────────────────────────────────────────────

def _extract_pdf(path: Path) -> Tuple[str, Optional[str]]:
    try:
        import pdfplumber  # type: ignore
        parts = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    parts.append(text)
        return "\n\n".join(parts), None
    except ImportError:
        return "", "pdfplumber not installed — run: pip install pdfplumber"
    except Exception as exc:
        return "", str(exc)


def _extract_docx(path: Path) -> Tuple[str, Optional[str]]:
    try:
        from docx import Document  # type: ignore
        doc = Document(path)
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text)
        # Also pull text from tables
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append("  |  ".join(cells))
        return "\n\n".join(paragraphs), None
    except ImportError:
        return "", "python-docx not installed — run: pip install python-docx"
    except Exception as exc:
        return "", str(exc)


def _extract_text_file(path: Path) -> Tuple[str, Optional[str]]:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
        return text, None
    except Exception as exc:
        return "", str(exc)


def extract_text(path: Path) -> Tuple[str, Optional[str]]:
    """Return (text, error_or_None) for any supported file type."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".docx":
        return _extract_docx(path)
    if ext in {".txt", ".md", ".rtf"}:
        return _extract_text_file(path)
    return "", f"unsupported type: {ext}"


# ── Cleaning ─────────────────────────────────────────────────────────────────

def clean_text(raw: str) -> str:
    """
    Normalise extracted text for downstream processing.

    Steps (in order):
      1. Unicode normalise to NFC (compose canonical equivalents)
      2. Replace non-printable / control characters with a space
         (keep newlines, tabs)
      3. Collapse runs of spaces on each line
      4. Collapse runs of blank lines to a single blank line
      5. Strip leading/trailing whitespace from the whole document
    """
    # 1. NFC normalise
    text = unicodedata.normalize("NFC", raw)

    # 2. Remove control characters except \n and \t
    text = re.sub(r"[^\S\n\t]", " ", text)  # non-newline whitespace → space
    # Remove non-printable chars that aren't whitespace
    text = "".join(
        ch if (unicodedata.category(ch)[0] != "C" or ch in "\n\t") else " "
        for ch in text
    )

    # 3. Collapse spaces within lines
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]

    # 4. Collapse consecutive blank lines
    cleaned_lines = []
    prev_blank = False
    for line in lines:
        is_blank = line == ""
        if is_blank and prev_blank:
            continue
        cleaned_lines.append(line)
        prev_blank = is_blank

    # 5. Strip
    return "\n".join(cleaned_lines).strip()


# ── Main ──────────────────────────────────────────────────────────────────────

def process_file(
    path: Path,
    out_dir: Path,
    force: bool = False,
) -> Tuple[bool, Optional[str]]:
    """
    Extract and clean one file, write to out_dir/<stem>.txt.
    Returns (success, error_or_None).
    """
    out_path = out_dir / (path.stem + ".txt")
    if out_path.exists() and not force:
        print(f"  SKIP  {path.name}  (already extracted — use --force to re-run)")
        return True, None

    raw_text, error = extract_text(path)
    if error:
        print(f"  ERROR {path.name}: {error}")
        return False, error

    cleaned = clean_text(raw_text)
    if not cleaned:
        print(f"  EMPTY {path.name}  (no text extracted)")
        return False, "no text extracted"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(cleaned, encoding="utf-8")
    print(f"  OK    {path.name}  →  {out_path.name}")
    return True, None


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Extract and clean text from documents."
    )
    parser.add_argument(
        "--source", type=Path, default=DATA_RAW,
        help="Folder containing source documents (default: data/raw)",
    )
    parser.add_argument(
        "--out", type=Path, default=DATA_PROCESSED,
        help="Output folder for .txt files (default: data/processed)",
    )
    parser.add_argument(
        "--force", action="store_true",
        help="Re-extract files that already have a .txt output",
    )
    args = parser.parse_args()

    if not args.source.exists():
        print(f"Source folder does not exist: {args.source}")
        sys.exit(1)

    files = [
        p for p in args.source.rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
    ]
    if not files:
        print(f"No supported files found in {args.source}")
        return

    print(f"Extracting {len(files)} files…")
    ok = errors = 0
    for path in sorted(files):
        success, _ = process_file(path, args.out, force=args.force)
        if success:
            ok += 1
        else:
            errors += 1

    print(f"\nDone: {ok} extracted, {errors} failed.")


if __name__ == "__main__":
    main()
