"""
knowledge/inventory.py
Document inventory script — run this first, before extracting or scrubbing.

Usage (PowerShell):
    python knowledge/inventory.py --source data/raw --out data/review/inventory.csv

What it does:
  - Scans a folder for PDF, DOCX, TXT and MD files
  - Records: filename, type, size, page/word count, extraction status
  - Detects exact duplicates (SHA-256 hash) and near-duplicates (Jaccard similarity)
  - Writes a CSV you can open in Excel before touching anything further
"""
import argparse
import csv
import hashlib
import re
import sys
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from config import DATA_RAW, DATA_REVIEW, SCRUB_SIMILARITY_THRESHOLD

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".rtf"}


# ── Hashing ─────────────────────────────────────────────────────────────────

def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


# ── Word / page count ────────────────────────────────────────────────────────

def _word_count_text(text: str) -> int:
    return len(text.split())


def _count_pdf(path: Path) -> Tuple[Optional[int], Optional[int], Optional[str]]:
    """Returns (pages, words, error)."""
    try:
        import pdfplumber  # type: ignore
        with pdfplumber.open(path) as pdf:
            pages = len(pdf.pages)
            words = sum(
                _word_count_text(p.extract_text() or "") for p in pdf.pages
            )
        return pages, words, None
    except ImportError:
        return None, None, "pdfplumber not installed"
    except Exception as exc:
        return None, None, str(exc)


def _count_docx(path: Path) -> Tuple[Optional[int], Optional[int], Optional[str]]:
    try:
        from docx import Document  # type: ignore
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        return None, _word_count_text(text), None
    except ImportError:
        return None, None, "python-docx not installed"
    except Exception as exc:
        return None, None, str(exc)


def _count_text_file(path: Path) -> Tuple[Optional[int], Optional[int], Optional[str]]:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
        return None, _word_count_text(text), None
    except Exception as exc:
        return None, None, str(exc)


def get_counts(path: Path) -> Tuple[Optional[int], Optional[int], Optional[str]]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _count_pdf(path)
    if ext == ".docx":
        return _count_docx(path)
    if ext in {".txt", ".md", ".rtf"}:
        return _count_text_file(path)
    return None, None, f"unsupported type: {ext}"


# ── Near-duplicate detection ──────────────────────────────────────────────────

def _ngrams(text: str, n: int = 3) -> set:
    """Character n-grams for Jaccard similarity."""
    text = re.sub(r"\s+", " ", text.lower())[:5000]  # cap at 5 000 chars for speed
    return {text[i:i+n] for i in range(len(text) - n + 1)}


def jaccard(a: str, b: str, n: int = 3) -> float:
    sa, sb = _ngrams(a, n), _ngrams(b, n)
    if not sa or not sb:
        return 0.0
    return len(sa & sb) / len(sa | sb)


def _extract_text_for_similarity(path: Path) -> str:
    """Best-effort text extraction for similarity comparison only."""
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            import pdfplumber  # type: ignore
            with pdfplumber.open(path) as pdf:
                return " ".join(p.extract_text() or "" for p in pdf.pages[:5])
        if ext == ".docx":
            from docx import Document  # type: ignore
            doc = Document(path)
            return " ".join(p.text for p in doc.paragraphs)
        return path.read_text(encoding="utf-8", errors="replace")
    except Exception:
        return ""


# ── Main ──────────────────────────────────────────────────────────────────────

def scan_folder(source_dir: Path) -> List[dict]:
    files = [
        p for p in source_dir.rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
    ]
    if not files:
        print(f"No supported files found in {source_dir}")
        return []

    print(f"Found {len(files)} files — scanning…")

    rows: List[dict] = []
    hashes: Dict[str, str] = {}  # hash -> first filename
    texts: Dict[str, str] = {}   # filename -> sample text

    for path in sorted(files):
        file_hash = sha256_file(path)
        exact_dup = hashes.get(file_hash)
        if not exact_dup:
            hashes[file_hash] = path.name

        pages, words, error = get_counts(path)
        sample_text = _extract_text_for_similarity(path)
        texts[path.name] = sample_text

        rows.append({
            "filename": path.name,
            "relative_path": str(path.relative_to(source_dir)),
            "extension": path.suffix.lower(),
            "size_bytes": path.stat().st_size,
            "pages": pages if pages is not None else "",
            "word_count": words if words is not None else "",
            "sha256": file_hash,
            "exact_duplicate_of": exact_dup if exact_dup and exact_dup != path.name else "",
            "near_duplicate_of": "",  # filled in below
            "extraction_error": error or "",
        })

    # Near-duplicate pass (O(n²) — fine for a personal library of ≤ a few hundred docs)
    names = [r["filename"] for r in rows]
    for i, row in enumerate(rows):
        if row["exact_duplicate_of"]:
            continue  # already flagged
        best_match = ""
        best_score = 0.0
        for j, other in enumerate(rows):
            if i == j:
                continue
            score = jaccard(texts.get(row["filename"], ""), texts.get(other["filename"], ""))
            if score > SCRUB_SIMILARITY_THRESHOLD and score > best_score:
                best_score = score
                best_match = f"{other['filename']} ({score:.2f})"
        row["near_duplicate_of"] = best_match

    return rows


def write_csv(rows: List[dict], out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "filename", "relative_path", "extension", "size_bytes",
        "pages", "word_count", "sha256",
        "exact_duplicate_of", "near_duplicate_of", "extraction_error",
    ]
    with out_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    print(f"Inventory written to {out_path}  ({len(rows)} files)")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Scan a folder of documents and write an inventory CSV."
    )
    parser.add_argument(
        "--source", type=Path, default=DATA_RAW,
        help="Folder containing source documents (default: data/raw)",
    )
    parser.add_argument(
        "--out", type=Path, default=DATA_REVIEW / "inventory.csv",
        help="Output CSV path (default: data/review/inventory.csv)",
    )
    args = parser.parse_args()

    if not args.source.exists():
        print(f"Source folder does not exist: {args.source}")
        sys.exit(1)

    rows = scan_folder(args.source)
    if rows:
        write_csv(rows, args.out)


if __name__ == "__main__":
    main()
